"""
Carrega as convenções coletivas de São Paulo no Supabase a partir da planilha
em `convencoes/`.

    python supabase/importar_convencoes.py [caminho-da-planilha]

São quatro, e o enquadramento sindical segue a ATIVIDADE PREPONDERANTE DA
EMPRESA (art. 511 §2º da CLT), não a função do empregado — uma empresa de
asseio que coloca um porteiro num cliente segue o SIEMACO-SP, e uma empresa de
prestação de serviços que faz o mesmo posto segue o SINDEEPRES. A diferença no
piso do porteiro é de R$ 131,03, o que vira uns R$ 400 no preço do posto.

É idempotente: rodar de novo atualiza cargos e benefícios pelo nome, sem
duplicar. Cargo já usado por alguma proposta mantém o id, então nada quebra.

O QUE NÃO É CARREGADO, DE PROPÓSITO

  · a tabela de "encargos sociais mínimos de 79,52%" que a CCT do asseio
    publica. Aquilo é o TOTAL — inclui 13º, férias, rescisão e cobertura, que
    aqui são os módulos 2.1, 3 e 4. Jogar no campo de encargos conta o mesmo
    custo duas vezes e infla o preço em cerca de 30%;
  · uniforme, EPI, arma, reciclagem e curso de formação: são Módulo 5
    (insumos), e vivem no catálogo de Equipamentos;
  · contribuição assistencial e sindical do EMPREGADO: é desconto na folha
    dele, não custo da empresa;
  · seguro de vida e assistência médica de valor não fixado na CCT ("prêmio da
    apólice"): entram quando a empresa souber o prêmio que paga. Inventar um
    número aqui seria pior do que deixar em branco.
"""
from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from dotenv import load_dotenv
from openpyxl import load_workbook

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
load_dotenv(os.path.join(RAIZ, ".env"))

import supa  # noqa: E402

PADRAO = os.path.join(RAIZ, "convencoes",
                      "Tabelas-CCT-SP-2026-Cargos-e-Beneficios_1.xlsx")


def _num(v, padrao=0.0) -> float:
    """Célula numérica vem como número e não se mexe. A limpeza de
    "R$ 1.837,40" vale só para texto: aplicada a um float, o replace do ponto
    transformaria 1837.4 em 18374."""
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace("R$", "").replace(".", "").replace(",", ".").strip())
    except (TypeError, ValueError):
        return padrao


def _txt(v) -> str:
    t = "" if v is None else str(v).strip()
    return "" if t in ("—", "-", "–") else t


# ════════════════════════════════════════════════════════════════════════════
# AS CONVENÇÕES
# ════════════════════════════════════════════════════════════════════════════
CONVENCOES = [
    {
        "chave": "SIEMACO",
        "aba_cargos": "1. SIEMACO - Cargos", "aba_benef": "2. SIEMACO - Beneficios",
        "cct": {
            "nome": "CCT SIEMACO-SP × SEAC-SP 2026/2027 — Asseio e conservação",
            "sindicato": "SIEMACO-SP (laboral) × SEAC-SP (patronal) · MTE SP003552/2026",
            "vigencia_inicio": "2026-01-01", "vigencia_fim": "2027-12-31",
            "uf": "SP", "cidade": "São Paulo", "insalubridade_base": "minimo",
            "salario_minimo": 1621.00, "piso_categoria": 1837.40, "horas_mensais": 220,
            "data_base": "2026-01-01", "reajuste_pct": 7.00,
            # A CCT do asseio AFASTA o adicional noturno depois das 05h no 12x36.
            "noturno_12x36_prorrogado": False,
            "intervalo_min": 30, "intervalo_adicional_pct": 50,
            "crts_pct": 0.40, "auxilio_creche_pct": 30.0,
            "cesta_teto_salario": 2720.86, "piso_jornada_reduzida_pct": 60,
            "ativo": True,
            "obs": "Asseio, conservação e limpeza urbana. 220 h/mês, já computados "
                   "os DSR. Jornada de 6h paga piso integral; abaixo de 4h, 60%. "
                   "A CCT publica encargos sociais mínimos de 79,5243% — esse número "
                   "é o TOTAL e NÃO vai no campo de encargos do sistema, que é só o "
                   "submódulo 2.2.",
        },
        "beneficios": [
            ("Vale-refeição", 21.80, "dia", 21.7, "pct_valor", 6.697),
            ("Vale-transporte", 10.40, "dia", 21.7, "pct_salario", 6.0),
            ("Cesta básica I (in natura)", 151.91, "mês", 0, "nenhum", 0),
            ("Cesta básica II (terceirizados)", 315.00, "mês", 0, "nenhum", 0),
            ("Assistência médica e odontológica", 37.09, "mês", 0, "nenhum", 0),
            ("Benefício Social Sindical", 16.75, "mês", 0, "nenhum", 0),
            ("Auxílio-creche", 486.30, "mês", 0, "nenhum", 0),
        ],
    },
    {
        "chave": "SESVESP",
        "aba_cargos": "3. SESVESP - Cargos", "aba_benef": "4. SESVESP - Beneficios",
        "cct": {
            "nome": "CCT SESVESP × FETRAVESP 2026/2027 — Vigilância patrimonial",
            "sindicato": "SESVESP (patronal) × FETRAVESP · MTE SP000195/2026",
            "vigencia_inicio": "2026-01-01", "vigencia_fim": "2027-12-31",
            "uf": "SP", "cidade": "São Paulo", "insalubridade_base": "minimo",
            "salario_minimo": 1621.00, "piso_categoria": 2271.74, "horas_mensais": 220,
            "data_base": "2026-01-01", "reajuste_pct": 5.75,
            "noturno_12x36_prorrogado": True,
            "intervalo_min": 60, "intervalo_adicional_pct": 50,
            "crts_pct": 0.0, "auxilio_creche_pct": 0.0,
            "cesta_teto_salario": 0.0, "piso_jornada_reduzida_pct": 60,
            "ativo": True,
            "obs": "Segurança privada patrimonial e pessoal. Periculosidade de 30% "
                   "devida ao vigilante (art. 193, II da CLT) — marque o adicional "
                   "P30 no posto. Hora extra convencional de 60%. Antecipação "
                   "salarial de 40% até o dia 20. Exige empresa autorizada pela "
                   "Polícia Federal, curso de formação e reciclagem bienal.",
        },
        "beneficios": [
            ("Vale-refeição", 42.00, "dia", 15.2, "pct_valor", 18.0),
            ("Vale-transporte", 10.40, "dia", 15.2, "pct_salario", 6.0),
            ("Cesta básica", 208.45, "mês", 0, "pct_valor", 5.0),
        ],
    },
    {
        "chave": "SINDEEPRES",
        "aba_cargos": "5. SINDEEPRES - Cargos", "aba_benef": "6. SINDEEPRES - Beneficios",
        "cct": {
            "nome": "CCT SINDEEPRES × SINDEPRESTEM 2026 — Prestação de serviços",
            "sindicato": "SINDEEPRES (laboral) × SINDEPRESTEM (patronal) · "
                         "MTE SP002405/2026 (geral) e SP002398/2026 (portaria)",
            "vigencia_inicio": "2026-01-01", "vigencia_fim": "2026-12-31",
            "uf": "SP", "cidade": "São Paulo", "insalubridade_base": "minimo",
            "salario_minimo": 1621.00, "piso_categoria": 1805.43, "horas_mensais": 220,
            "data_base": "2026-01-01", "reajuste_pct": 6.25,
            "noturno_12x36_prorrogado": True,
            "intervalo_min": 60, "intervalo_adicional_pct": 50,
            "crts_pct": 0.0, "auxilio_creche_pct": 0.0,
            "cesta_teto_salario": 0.0, "piso_jornada_reduzida_pct": 60,
            "ativo": True,
            "obs": "Prestação de serviços a terceiros e colocação de mão de obra. "
                   "VIGÊNCIA ANUAL: 01/01/2026 a 31/12/2026 — renegocia todo ano, "
                   "diferente das outras. Reajuste de 6,25% até R$ 7.380,07 e 5,50% "
                   "acima. Dois aditivos com pisos diferentes: GERAL e PORTARIA. "
                   "Não prevê assistência médica nem auxílio-creche, ao contrário "
                   "do SIEMACO-SP. O prêmio de boa permanência de R$ 110,00 é "
                   "mensal e habitual: verifique com o jurídico se não tem "
                   "natureza salarial, porque aí sofre encargo.",
        },
        "beneficios": [
            ("Vale-refeição (aditivo geral)", 24.80, "dia", 21.7, "nenhum", 0),
            ("Vale-refeição (aditivo portaria)", 26.03, "dia", 21.7, "nenhum", 0),
            ("Vale-transporte", 10.40, "dia", 21.7, "pct_salario", 6.0),
            ("Cesta básica (aditivo geral)", 174.10, "mês", 0, "nenhum", 0),
            ("Cesta básica (aditivo portaria)", 205.91, "mês", 0, "nenhum", 0),
            ("Prêmio de boa permanência", 110.00, "mês", 0, "nenhum", 0),
            ("Assistência odontológica", 28.31, "mês", 0, "nenhum", 0),
            ("PLR mensalizada", 29.30, "mês", 0, "nenhum", 0),
        ],
    },
    {
        "chave": "BOMBEIROS",
        "aba_cargos": "7. BOMBEIROS - Cargos", "aba_benef": "8. BOMBEIROS - Beneficios",
        "cct": {
            "nome": "CCT SINDIBOMBEIROS × SINDEPRESTEM 2025/2026 — Bombeiro civil",
            "sindicato": "SINDIBOMBEIROS (laboral) × SINDEPRESTEM (patronal) · "
                         "MTE SP012085/2025",
            "vigencia_inicio": "2025-09-01", "vigencia_fim": "2026-08-31",
            "uf": "SP", "cidade": "São Paulo", "insalubridade_base": "minimo",
            "salario_minimo": 1621.00, "piso_categoria": 2784.30, "horas_mensais": 180,
            # ⚠ data-base de SETEMBRO, não de janeiro como as outras três.
            "data_base": "2025-09-01", "reajuste_pct": 5.50,
            "noturno_12x36_prorrogado": True,
            "intervalo_min": 60, "intervalo_adicional_pct": 50,
            "crts_pct": 0.0, "auxilio_creche_pct": 0.0,
            "cesta_teto_salario": 0.0, "piso_jornada_reduzida_pct": 60,
            "ativo": True,
            "obs": "Bombeiro profissional civil. ATENÇÃO À DATA-BASE: 1º de "
                   "SETEMBRO, diferente das outras três convenções — um contrato "
                   "que misture bombeiro com asseio tem DUAS datas de repactuação "
                   "no ano. Divisor de 180 h/mês para o operacional (Lei "
                   "11.901/2009) e 220 h para chefia e técnicos: o divisor está "
                   "gravado em cada cargo. Escala 12x36 obrigatória (cláusula 58ª). "
                   "Periculosidade de 30% sobre o salário. Hora extra em folga: "
                   "100%. PLR de R$ 1.766,53 ao ano.",
        },
        "beneficios": [
            ("Vale-refeição", 34.26, "dia", 15.2, "nenhum", 0),
            ("Vale-transporte", 10.40, "dia", 15.2, "pct_salario", 6.0),
            ("Cesta básica", 193.62, "mês", 0, "nenhum", 0),
            ("PLR mensalizada", 147.21, "mês", 0, "nenhum", 0),
            ("Assistência odontológica", 23.00, "mês", 0, "pct_valor", 100.0),
        ],
    },
]

# Módulo 5 — o que a planilha traz de insumo, por categoria.
EQUIPAMENTOS = [
    ("Reciclagem bienal de vigilante", "Lei 7.102/83 — exigência da PF", 67.50, "Mensal"),
    ("Arma, munição e colete balístico", "Depreciação, porte e estande", 140.00, "Mensal"),
    ("Uniforme de vigilante", "Especificação da CCT — 2 conjuntos", 125.00, "Mensal"),
    ("Uniforme de asseio", "2 conjuntos e agasalho — CCT", 77.50, "Mensal"),
    ("EPI por função", "NR-06 — fornecimento gratuito", 55.00, "Mensal"),
    ("Reciclagem anual de bombeiro civil", "Cláusula 41ª — custeada pela empresa", 90.00, "Mensal"),
    ("Credenciamento CBPMESP", "Decreto Estadual 69.118/2024", 45.00, "Mensal"),
]


# ════════════════════════════════════════════════════════════════════════════
# LEITURA DAS ABAS
# ════════════════════════════════════════════════════════════════════════════
def cargos_siemaco(ws):
    """Função | piso | valor-hora | insalubridade usual | periculosidade usual | obs"""
    saida = []
    for r in ws.iter_rows(min_row=6, max_row=26, values_only=True):
        nome, salario = _txt(r[0]), _num(r[1])
        if not nome or salario <= 0 or nome.isupper():
            continue
        partes = []
        if _txt(r[3]): partes.append("Insalubridade: " + _txt(r[3]))
        if _txt(r[4]): partes.append("Periculosidade: " + _txt(r[4]))
        if _txt(r[5]): partes.append(_txt(r[5]))
        saida.append({"nome": nome, "salario": round(salario, 2),
                      "obs": " · ".join(partes) or None})
    return saida


def cargos_sesvesp(ws):
    """Função | piso | gratificação | valor da gratificação | periculosidade |
    remuneração mínima | obs.

    O salário gravado é o piso MAIS a gratificação de função, que é fixa do
    cargo. A periculosidade fica de fora: ela é adicional do posto (P30), e
    somá-la aqui a contaria duas vezes."""
    saida, grupo = [], ""
    for r in ws.iter_rows(min_row=6, max_row=28, values_only=True):
        nome = _txt(r[0])
        if not nome:
            continue
        if nome.upper().startswith("GRUPO"):
            grupo = nome.split("—")[-1].strip().title()
            continue
        piso = _num(r[1])
        if piso <= 0:
            continue
        grat = _num(r[3])
        partes = [p for p in (grupo,) if p]
        if _txt(r[2]): partes.append("Gratificação de função: " + _txt(r[2]))
        if _num(r[4]) > 0:
            partes.append("Periculosidade 30%% = R$ %.2f — marque P30 no posto" % _num(r[4]))
        if len(r) > 6 and _txt(r[6]): partes.append(_txt(r[6]))
        saida.append({"nome": nome, "salario": round(piso + grat, 2),
                      "obs": " · ".join(partes) or None})
    return saida


def cargos_sindeepres(ws):
    """Função | aditivo | piso | valor-hora | obs. Dois aditivos com pisos
    diferentes para a mesma função — o aditivo entra no nome, senão viram dois
    cargos indistinguíveis na lista."""
    saida = []
    for r in ws.iter_rows(min_row=6, max_row=34, values_only=True):
        nome, aditivo, salario = _txt(r[0]), _txt(r[1]), _num(r[2])
        if not nome or salario <= 0 or nome.isupper():
            continue
        rotulo = "%s (%s)" % (nome, aditivo.lower()) if aditivo else nome
        saida.append({"nome": rotulo, "salario": round(salario, 2),
                      "aditivo": aditivo or None,
                      "obs": _txt(r[4]) or None})
    return saida


def cargos_bombeiros(ws):
    """Função | piso | gratificação | valor da gratificação | periculosidade |
    remuneração mínima | divisor | valor-hora.

    O divisor muda dentro da convenção — 180 h no operacional, 220 na chefia — e
    vai para a coluna `horas_mensais` do cargo. O salário gravado é o piso mais
    a gratificação de função, que é fixa por cargo; a periculosidade continua
    sendo adicional do posto."""
    saida, grupo = [], ""
    for r in ws.iter_rows(min_row=6, max_row=40, values_only=True):
        nome = _txt(r[0])
        if not nome:
            continue
        if nome.isupper() and _num(r[1]) <= 0:
            grupo = nome.title()
            continue
        piso = _num(r[1])
        if piso <= 0:
            continue
        grat = _num(r[3])
        divisor = 180.0 if "180" in _txt(r[6]) else 220.0
        partes = [p for p in (grupo,) if p]
        if _txt(r[2]): partes.append("Gratificação de função: " + _txt(r[2]))
        if _num(r[4]) > 0:
            partes.append("Periculosidade 30%% = R$ %.2f — marque P30 no posto" % _num(r[4]))
        partes.append("Divisor de %d h/mês" % divisor)
        saida.append({"nome": nome, "salario": round(piso + grat, 2),
                      "horas_mensais": divisor,
                      "obs": " · ".join(partes) or None})
    return saida


LEITORES = {"SIEMACO": cargos_siemaco, "SESVESP": cargos_sesvesp,
            "SINDEEPRES": cargos_sindeepres, "BOMBEIROS": cargos_bombeiros}


# ════════════════════════════════════════════════════════════════════════════
def sincronizar(sb, dados, cargos, beneficios):
    achadas = sb.select("ccts", nome="eq.%s" % dados["nome"])
    if achadas:
        cct = achadas[0]
        sb.update("ccts", dados, id="eq.%s" % cct["id"])
        acao = "atualizada"
    else:
        cct = sb.insert("ccts", dados)
        acao = "criada"
    cid = cct["id"]

    exist = {c["nome"]: c for c in sb.select("cct_cargos", **{"cctId": "eq.%s" % cid})}
    novos = atual = 0
    for c in cargos:
        linha = dict(c)
        if c["nome"] in exist:
            sb.update("cct_cargos", linha, id="eq.%s" % exist[c["nome"]]["id"])
            atual += 1
        else:
            sb.insert("cct_cargos", dict(linha, **{"cctId": cid}), retornar=False)
            novos += 1

    exist_b = {b["nome"]: b for b in sb.select("cct_beneficios", **{"cctId": "eq.%s" % cid})}
    nb = ab = 0
    for nome, valor, unid, dias, tipo, pct in beneficios:
        linha = {"nome": nome, "valor": valor, "unid": unid,
                 "dias_mes": dias or 22, "desconto_tipo": tipo, "desconto_pct": pct}
        if nome in exist_b:
            sb.update("cct_beneficios", linha, id="eq.%s" % exist_b[nome]["id"])
            ab += 1
        else:
            sb.insert("cct_beneficios", dict(linha, **{"cctId": cid}), retornar=False)
            nb += 1
    print("  %-11s convenção %-10s cargos %2d novos / %2d atualizados · "
          "benefícios %d novos / %d atualizados"
          % (dados["nome"][:11], acao, novos, atual, nb, ab))
    return cid


def main():
    caminho = sys.argv[1] if len(sys.argv) > 1 else PADRAO
    if not os.path.isfile(caminho):
        print("planilha não encontrada: %s" % caminho)
        return 1
    wb = load_workbook(caminho, data_only=True)
    sb = supa.Supabase(supa.service_key())

    print("\nCONVENÇÕES")
    for conf in CONVENCOES:
        cargos = LEITORES[conf["chave"]](wb[conf["aba_cargos"]])
        sincronizar(sb, conf["cct"], cargos, conf["beneficios"])

    print("\nMÓDULO 5 — insumos")
    existentes = {e["nome"] for e in sb.select("equipamentos")}
    n = 0
    for nome, marca, valor, tipo in EQUIPAMENTOS:
        if nome not in existentes:
            sb.insert("equipamentos", {"nome": nome, "marca_modelo": marca,
                                       "valor": valor, "tipo": tipo}, retornar=False)
            n += 1
    print("  %d itens criados" % n)
    return 0


if __name__ == "__main__":
    sys.exit(main())
